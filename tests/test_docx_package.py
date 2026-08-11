from openpyxl import Workbook
from docx import Document

from app.export.docx_package import write_docx_report
from app.pipeline.integrated import run_integrated_period

HEADERS = ["Pauta de transmisión", "Estado", "Tiempo Fiscal", "Canal Base", "Orden", "Fecha", "Dependencia CAM. SEN.", "Clave", "Campaña", "Versión"]


def test_write_docx_report_creates_readable_document(tmp_path) -> None:
    source = tmp_path / "rtc.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    sheet.append(["Pauta", "Programada", "00:00:30", "AM", "O1", "11/08/2026", "CAM. SEN.", "C1", "Campaña", "V1"])
    workbook.save(source)

    result = run_integrated_period([source], "2026-Q2")
    output = tmp_path / "Informe_Ejecutivo.docx"
    write_docx_report(result, output)
    document = Document(output)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Informe Ejecutivo — SIAP-RTC" in text
    assert "Consideraciones metodológicas" in text
